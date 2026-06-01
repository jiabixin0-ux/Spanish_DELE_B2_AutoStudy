from __future__ import annotations

import re
import sys
from pathlib import Path


PROJECT_ROOT = Path(__file__).resolve().parents[1]
VENDOR_DIR = PROJECT_ROOT / "vendor"
if VENDOR_DIR.exists() and str(VENDOR_DIR) not in sys.path:
    sys.path.insert(0, str(VENDOR_DIR))

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


MARKDOWN_PATH = PROJECT_ROOT / "output" / "daily_lesson.md"
DOCX_PATH = PROJECT_ROOT / "output" / "daily_lesson.docx"
MARKDOWN_DISPLAY_PATH = Path("output") / "daily_lesson.md"
DOCX_DISPLAY_PATH = Path("output") / "daily_lesson.docx"


BODY_FONT = "Arial"
EAST_ASIA_FONT = "PingFang SC"
ACCENT_BLUE = "1F4D78"
SOFT_BLUE = "D9EAF7"
SOFT_GRAY = "F4F6F9"
SPANISH_BLUE = "17365D"
CHINESE_GRAY = "555555"


def set_run_font(run, size_pt: float = 11.0, bold: bool = False) -> None:
    run.font.name = BODY_FONT
    run.font.size = Pt(size_pt)
    run.bold = bold
    run._element.rPr.rFonts.set(qn("w:eastAsia"), EAST_ASIA_FONT)


def set_style_font(style, size_pt: float, bold: bool = False, color: str | None = None) -> None:
    style.font.name = BODY_FONT
    style.font.size = Pt(size_pt)
    style.font.bold = bold
    if color:
        style.font.color.rgb = RGBColor.from_string(color)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), EAST_ASIA_FONT)


def set_paragraph_shading(paragraph, fill: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shading = p_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        p_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def add_bottom_border(paragraph, color: str = "B7C4D6") -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    border = p_pr.find(qn("w:pBdr"))
    if border is None:
        border = OxmlElement("w:pBdr")
        p_pr.append(border)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), color)
    border.append(bottom)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_width(cell, width_twips: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    width = tc_pr.find(qn("w:tcW"))
    if width is None:
        width = OxmlElement("w:tcW")
        tc_pr.append(width)
    width.set(qn("w:w"), str(width_twips))
    width.set(qn("w:type"), "dxa")


def set_cell_margins(cell, margin_twips: int = 100) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.find(qn("w:tcMar"))
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for edge in ("top", "left", "bottom", "right"):
        node = margins.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            margins.append(node)
        node.set(qn("w:w"), str(margin_twips))
        node.set(qn("w:type"), "dxa")


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.85)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    styles = document.styles
    normal = styles["Normal"]
    set_style_font(normal, 11.5)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.2

    for style_name, size, space_before, space_after, color in (
        ("Heading 1", 17, 18, 10, ACCENT_BLUE),
        ("Heading 2", 13.5, 12, 6, ACCENT_BLUE),
    ):
        style = styles[style_name]
        set_style_font(style, size, bold=True, color=color)
        style.paragraph_format.space_before = Pt(space_before)
        style.paragraph_format.space_after = Pt(space_after)
        style.paragraph_format.line_spacing = 1.15

    if "Lesson Title" not in styles:
        title_style = styles.add_style("Lesson Title", WD_STYLE_TYPE.PARAGRAPH)
    else:
        title_style = styles["Lesson Title"]
    set_style_font(title_style, 24, bold=True, color=ACCENT_BLUE)
    title_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_style.paragraph_format.space_after = Pt(18)
    title_style.paragraph_format.line_spacing = 1.15

    for style_name, size, color, italic, before, after in (
        ("Spanish Example", 11.5, SPANISH_BLUE, False, 0, 6),
        ("Chinese Translation", 11, CHINESE_GRAY, False, 0, 8),
        ("Label Line", 11.2, None, False, 0, 6),
        ("Instruction", 11, "333333", False, 0, 8),
    ):
        style = styles.add_style(style_name, WD_STYLE_TYPE.PARAGRAPH) if style_name not in styles else styles[style_name]
        set_style_font(style, size, color=color)
        style.font.italic = italic
        style.paragraph_format.left_indent = Inches(0.12)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.2


def strip_markdown_emphasis(text: str) -> str:
    text = text.replace("`", "")
    text = re.sub(r"\*\*(.*?)\*\*", r"\1", text)
    text = re.sub(r"\*(.*?)\*", r"\1", text)
    return text.strip()


def add_text_paragraph(document: Document, text: str, style: str | None = None) -> None:
    paragraph = document.add_paragraph(style=style)
    paragraph.paragraph_format.space_after = Pt(8)
    for index, part in enumerate(split_spanish_chinese(text)):
        if index:
            paragraph.add_run().add_break()
        run = paragraph.add_run(part)
        set_run_font(run)


def add_list_paragraph(document: Document, text: str, numbered: bool = False) -> None:
    paragraph = document.add_paragraph(style="List Number" if numbered else "List Bullet")
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.left_indent = Inches(0.25)
    for index, part in enumerate(split_spanish_chinese(strip_markdown_emphasis(text))):
        if index:
            paragraph.add_run().add_break()
        run = paragraph.add_run(part)
        set_run_font(run)


def split_spanish_chinese(text: str) -> list[str]:
    text = strip_markdown_emphasis(text)
    if "：" in text:
        left, right = text.split("：", 1)
        if re.search(r"[A-Za-zÁÉÍÓÚÜÑáéíóúüñ]", left) and re.search(r"[\u4e00-\u9fff]", right):
            return [left.strip() + "：", right.strip()]
    if " - " in text:
        left, right = text.split(" - ", 1)
        if re.search(r"[A-Za-zÁÉÍÓÚÜÑáéíóúüñ]", left) and re.search(r"[\u4e00-\u9fff]", right):
            return [left.strip(), right.strip()]
    return [text]


def add_blockquote(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.25)
    paragraph.paragraph_format.right_indent = Inches(0.1)
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(8)
    p_pr = paragraph._p.get_or_add_pPr()
    border = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "12")
    left.set(qn("w:space"), "6")
    left.set(qn("w:color"), "B7C4D6")
    border.append(left)
    p_pr.append(border)
    run = paragraph.add_run(strip_markdown_emphasis(text))
    set_run_font(run, size_pt=10.5)
    run.italic = True


def add_title(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(style="Lesson Title")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(strip_markdown_emphasis(text))
    set_run_font(run, size_pt=24, bold=True)
    run.font.color.rgb = RGBColor.from_string(ACCENT_BLUE)
    add_bottom_border(paragraph)


def add_heading_one(document: Document, text: str) -> None:
    paragraph = document.add_heading(strip_markdown_emphasis(text), level=1)
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.space_before = Pt(18)
    paragraph.paragraph_format.space_after = Pt(10)
    set_paragraph_shading(paragraph, SOFT_BLUE)
    add_bottom_border(paragraph, color="9CBAD6")


def add_heading_two(document: Document, text: str) -> None:
    paragraph = document.add_heading(strip_markdown_emphasis(text), level=2)
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.space_before = Pt(12)
    paragraph.paragraph_format.space_after = Pt(6)


def add_label_paragraph(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(style="Label Line")
    label_match = re.match(r"^(【[^】]+】)(.*)$", strip_markdown_emphasis(text))
    if label_match:
        label, rest = label_match.groups()
        label_run = paragraph.add_run(label)
        set_run_font(label_run, size_pt=11.2, bold=True)
        label_run.font.color.rgb = RGBColor.from_string(ACCENT_BLUE)
        if rest.strip():
            rest_run = paragraph.add_run(rest.strip())
            set_run_font(rest_run, size_pt=11.2)
    else:
        run = paragraph.add_run(strip_markdown_emphasis(text))
        set_run_font(run, bold=True)


def add_example_label(document: Document, label: str) -> None:
    paragraph = document.add_paragraph(style="Label Line")
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run(label)
    set_run_font(run, size_pt=11, bold=True)
    run.font.color.rgb = RGBColor.from_string(ACCENT_BLUE)


def add_example_text(document: Document, text: str, chinese: bool = False) -> None:
    text = text.removeprefix("> ").strip()
    paragraph = document.add_paragraph(style="Chinese Translation" if chinese else "Spanish Example")
    paragraph.paragraph_format.left_indent = Inches(0.28)
    paragraph.paragraph_format.space_after = Pt(8 if chinese else 4)
    if not chinese:
        set_paragraph_shading(paragraph, SOFT_GRAY)
    run = paragraph.add_run(strip_markdown_emphasis(text))
    set_run_font(run, size_pt=11 if chinese else 11.5)
    if chinese:
        run.font.color.rgb = RGBColor.from_string(CHINESE_GRAY)
    else:
        run.font.color.rgb = RGBColor.from_string(SPANISH_BLUE)


def is_table_line(line: str) -> bool:
    stripped = line.strip()
    return stripped.startswith("|") and stripped.endswith("|") and stripped.count("|") >= 2


def is_table_separator(line: str) -> bool:
    cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", cell or "") for cell in cells)


def split_table_row(line: str) -> list[str]:
    return [strip_markdown_emphasis(cell.strip()) for cell in line.strip().strip("|").split("|")]


def add_markdown_table(document: Document, table_lines: list[str]) -> None:
    rows = [split_table_row(line) for line in table_lines if not is_table_separator(line)]
    if not rows:
        return
    column_count = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=column_count)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    widths = [1500, 1300, 3600, 3200]

    for row_index, row in enumerate(rows):
        for col_index in range(column_count):
            cell = table.cell(row_index, col_index)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell, 110)
            if col_index < len(widths):
                set_cell_width(cell, widths[col_index])
            if row_index == 0:
                set_cell_shading(cell, SOFT_BLUE)
            text = row[col_index] if col_index < len(row) else ""
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(2)
            paragraph.paragraph_format.line_spacing = 1.12
            run = paragraph.add_run(text)
            set_run_font(run, size_pt=9.3 if row_index else 9.8, bold=row_index == 0)
            if row_index == 0:
                run.font.color.rgb = RGBColor.from_string(ACCENT_BLUE)

    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(8)


def markdown_to_docx(markdown: str, output_path: Path) -> None:
    document = Document()
    configure_document(document)
    active_example_style: str | None = None

    lines = markdown.splitlines()
    index = 0
    while index < len(lines):
        raw_line = lines[index]
        line = raw_line.strip()
        index += 1
        if not line:
            continue
        if is_table_line(line):
            table_lines = [line]
            while index < len(lines) and is_table_line(lines[index].strip()):
                table_lines.append(lines[index].strip())
                index += 1
            add_markdown_table(document, table_lines)
            continue
        if line.startswith("# "):
            active_example_style = None
            add_title(document, line[2:])
            continue
        if line.startswith("## "):
            active_example_style = None
            add_heading_one(document, line[3:])
            continue
        if line.startswith("### "):
            active_example_style = None
            add_heading_two(document, line[4:])
            continue
        if line in {"西语：", "中文：", "西语原文：", "中文翻译："}:
            add_example_label(document, line)
            active_example_style = "chinese" if line.startswith("中文") else "spanish"
            continue
        if line.startswith("【"):
            add_label_paragraph(document, line)
            if line.startswith(("【西语原文】", "【例句】")):
                active_example_style = "spanish"
            elif line.startswith(("【中文翻译】", "【翻译】")):
                active_example_style = "chinese"
            else:
                active_example_style = None
            continue
        if active_example_style:
            add_example_text(document, line, chinese=active_example_style == "chinese")
            continue
        if line.startswith("- "):
            add_list_paragraph(document, line[2:], numbered=False)
            continue
        numbered_match = re.match(r"^(\d+)\.\s+(.*)$", line)
        if numbered_match:
            add_text_paragraph(document, f"{numbered_match.group(1)}. {numbered_match.group(2)}", style="Instruction")
            continue
        if line.startswith("> "):
            add_example_text(document, line[2:], chinese=False)
            continue
        add_text_paragraph(document, line, style="Instruction")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)


def main() -> int:
    if not MARKDOWN_PATH.exists():
        raise SystemExit(f"缺少 {MARKDOWN_DISPLAY_PATH}，请先运行 scripts/generate_daily_lesson.py")
    markdown = MARKDOWN_PATH.read_text(encoding="utf-8")
    markdown_to_docx(markdown, DOCX_PATH)
    print(f"已生成 Word 附件：{DOCX_DISPLAY_PATH}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
